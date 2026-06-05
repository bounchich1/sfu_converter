from docx import Document

from sfu_converter.infrastructure.main_inscription import graph_text, render


def test_main_inscription_renders_all_seventeen_graphs_and_fields():
    doc = Document()

    table = render(
        doc,
        "form_1",
        fields={"1": "Пояснительная записка", "2": "КП-01.02.03 ПЗ"},
    )

    assert len(table.rows) == 17
    assert table.cell(0, 0).text == "Форма form_1"
    assert graph_text(table, 1) == "Пояснительная записка"
    assert graph_text(table, 2) == "КП-01.02.03 ПЗ"
    assert graph_text(table, 17) == ""


def test_main_inscription_can_place_page_field_in_graph_seven():
    doc = Document()

    table = render(doc, "form_3", fields={}, page_number_in_graph_7=True)

    assert " PAGE " in table.rows[6].cells[1]._tc.xml
